"""Build an Excel recruitment tracker (candidate log + live dashboard) from the CSV.

The workbook uses formulas (COUNTIF/COUNTIFS/AVERAGE), dropdown validation and
conditional formatting, so it keeps updating when a recruiter edits a row.

Usage:
    python -m src.build_tracker --input data/sample_candidates.csv --out output/recruitment_tracker.xlsx
"""
import argparse
import csv
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .config import OUTCOMES, RECRUITERS, ROLES, SCREEN_SLA_DAYS, SOURCES, STAGES

FONT = "Arial"
NAVY, TEAL, LIGHT = "12284C", "0A7EA4", "E8F1F6"
MAX_ROW = 500  # formulas and dropdowns cover rows 2..500 so new candidates are picked up

HEADERS = [
    "Candidate ID", "Name", "Email", "Role", "Source", "Recruiter",
    "Applied Date", "Screening Date", "Interview Date", "Offer Date",
    "Stage", "Outcome", "Notes",
    "Stage #", "Days to Screen", "Days to Offer",
]
INPUT_COLS = 13  # columns A..M are typed by the recruiter; N..P are formulas

thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def hdr(cell, fill=NAVY):
    cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
    cell.fill = PatternFill("solid", start_color=fill)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER


def to_date(s):
    return date.fromisoformat(s) if s else None


def read_csv(path):
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def add_lists_sheet(wb):
    ws = wb.create_sheet("Lists")
    cols = {"Stage": STAGES, "Outcome": OUTCOMES, "Role": ROLES, "Source": SOURCES, "Recruiter": RECRUITERS}
    for c, (name, vals) in enumerate(cols.items(), start=1):
        hdr(ws.cell(row=1, column=c, value=name))
        for r, v in enumerate(vals, start=2):
            ws.cell(row=r, column=c, value=v).font = Font(name=FONT, size=10)
        ws.column_dimensions[get_column_letter(c)].width = 28
    return {name: f"Lists!${get_column_letter(c)}$2:${get_column_letter(c)}${len(vals) + 1}"
            for c, (name, vals) in enumerate(cols.items(), start=1)}


def add_candidates_sheet(wb, rows, ranges):
    ws = wb.create_sheet("Candidates")
    for c, h in enumerate(HEADERS, start=1):
        hdr(ws.cell(row=1, column=c, value=h), fill=NAVY if c <= INPUT_COLS else TEAL)
    ws.row_dimensions[1].height = 30

    for r, row in enumerate(rows, start=2):
        vals = [row["candidate_id"], row["name"], row["email"], row["role"], row["source"], row["recruiter"],
                to_date(row["applied_date"]), to_date(row["screening_date"]),
                to_date(row["interview_date"]), to_date(row["offer_date"]),
                row["stage"], row["outcome"], row["notes"] or None]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = Font(name=FONT, size=10, color="0000FF")  # blue = typed input
            cell.border = BORDER
            if 7 <= c <= 10:
                cell.number_format = "dd-mmm-yyyy"

    for r in range(2, MAX_ROW + 1):
        ws.cell(row=r, column=14, value=f'=IFERROR(MATCH(K{r},{ranges["Stage"]},0),"")')
        ws.cell(row=r, column=15, value=f'=IF(OR(G{r}="",H{r}=""),"",H{r}-G{r})')
        ws.cell(row=r, column=16, value=f'=IF(OR(G{r}="",J{r}=""),"",J{r}-G{r})')
        for c in (14, 15, 16):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=FONT, size=10, color="000000")  # black = formula
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center")

    for col, key in (("D", "Role"), ("E", "Source"), ("F", "Recruiter"), ("K", "Stage"), ("L", "Outcome")):
        dv = DataValidation(type="list", formula1=f"={ranges[key]}", allow_blank=True,
                            showErrorMessage=True, errorTitle="Invalid value",
                            error="Pick a value from the dropdown list (see the Lists sheet).")
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{MAX_ROW}")
    dv_date = DataValidation(type="date", operator="greaterThan", formula1="43831", allow_blank=True,
                             showErrorMessage=True, errorTitle="Invalid date", error="Enter a valid date.")
    ws.add_data_validation(dv_date)
    dv_date.add(f"G2:J{MAX_ROW}")

    rng = f"L2:L{MAX_ROW}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Joined"'],
                                  fill=PatternFill("solid", bgColor="C6EFCE"), font=Font(color="006100")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Rejected"'],
                                  fill=PatternFill("solid", bgColor="FFC7CE"), font=Font(color="9C0006")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Withdrawn"'],
                                  fill=PatternFill("solid", bgColor="E7E6E6"), font=Font(color="595959")))
    ws.conditional_formatting.add(f"O2:O{MAX_ROW}", FormulaRule(
        formula=[f'AND(ISNUMBER(O2),O2>Dashboard!$B$4)'],
        fill=PatternFill("solid", bgColor="FFEB9C"), font=Font(color="9C5700")))

    for col, w in zip("ABCDEFGHIJKLMNOP", [12, 20, 34, 27, 19, 14, 14, 14, 14, 14, 13, 13, 22, 9, 11, 11]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:P{MAX_ROW}"
    return ws


def add_dashboard(wb):
    ws = wb.create_sheet("Dashboard", 0)
    ws.sheet_view.showGridLines = False
    C = "Candidates"
    rng = lambda col: f"{C}!${col}$2:${col}${MAX_ROW}"

    ws["A1"] = "Recruitment Dashboard"
    ws["A1"].font = Font(name=FONT, bold=True, size=16, color=NAVY)
    ws["A2"] = "All figures are formulas linked to the Candidates sheet. Data is synthetic."
    ws["A2"].font = Font(name=FONT, italic=True, size=9, color="595959")

    ws["A4"] = "Screening SLA (days)"
    ws["B4"] = SCREEN_SLA_DAYS
    ws["A4"].font = Font(name=FONT, bold=True, size=10)
    ws["B4"].font = Font(name=FONT, size=10, color="0000FF")  # blue = editable assumption
    ws["B4"].fill = PatternFill("solid", start_color="FFFF00")
    ws["C4"] = "Edit B4: screening delays longer than this are highlighted on the Candidates sheet."
    ws["C4"].font = Font(name=FONT, italic=True, size=9, color="595959")

    def kpi(row, label, formula, fmt="0"):
        ws.cell(row=row, column=1, value=label).font = Font(name=FONT, bold=True, size=10)
        c = ws.cell(row=row, column=2, value=formula)
        c.font = Font(name=FONT, size=10)
        c.number_format = fmt
        c.alignment = Alignment(horizontal="right")

    kpi(5, "Total candidates", f"=COUNTA({rng('A')})")
    kpi(6, "Active (In Progress)", f'=COUNTIF({rng("L")},"In Progress")')
    kpi(7, "Avg days to screen", f"=IFERROR(AVERAGE({rng('O')}),0)", "0.0")
    kpi(8, "Avg days to offer", f"=IFERROR(AVERAGE({rng('P')}),0)", "0.0")
    kpi(9, "% screened within SLA", f'=IFERROR(COUNTIFS({rng("O")},"<="&B4)/COUNT({rng("O")}),0)', "0%")

    # Funnel
    r0 = 11
    ws.cell(row=r0, column=1, value="Hiring funnel (reached at least this stage)").font = Font(name=FONT, bold=True, size=11, color=NAVY)
    for c, h in enumerate(["Stage", "Stage #", "Candidates", "Step conversion", "Overall conversion"], start=1):
        hdr(ws.cell(row=r0 + 1, column=c, value=h))
    for i, stage in enumerate(STAGES):
        r = r0 + 2 + i
        ws.cell(row=r, column=1, value=stage)
        ws.cell(row=r, column=2, value=i + 1)
        ws.cell(row=r, column=3, value=f'=COUNTIF({rng("N")},">="&B{r})')
        ws.cell(row=r, column=4, value=(f"=IFERROR(C{r}/C{r-1},0)" if i else None))
        ws.cell(row=r, column=5, value=f"=IFERROR(C{r}/C${r0+2},0)")
        for c in range(1, 6):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=FONT, size=10)
            cell.border = BORDER
            if c in (4, 5):
                cell.number_format = "0%"
    funnel_first, funnel_last = r0 + 2, r0 + 1 + len(STAGES)

    # Source effectiveness
    s0 = funnel_last + 3
    ws.cell(row=s0, column=1, value="Sourcing channel effectiveness").font = Font(name=FONT, bold=True, size=11, color=NAVY)
    for c, h in enumerate(["Source", "Applicants", "Interviewed+", "Offered+", "Joined", "Interview rate", "Offer rate"], start=1):
        hdr(ws.cell(row=s0 + 1, column=c, value=h))
    for i, src in enumerate(SOURCES):
        r = s0 + 2 + i
        ws.cell(row=r, column=1, value=src)
        ws.cell(row=r, column=2, value=f"=COUNTIF({rng('E')},A{r})")
        ws.cell(row=r, column=3, value=f'=COUNTIFS({rng("E")},A{r},{rng("N")},">=3")')
        ws.cell(row=r, column=4, value=f'=COUNTIFS({rng("E")},A{r},{rng("N")},">=4")')
        ws.cell(row=r, column=5, value=f'=COUNTIFS({rng("E")},A{r},{rng("L")},"Joined")')
        ws.cell(row=r, column=6, value=f"=IFERROR(C{r}/B{r},0)")
        ws.cell(row=r, column=7, value=f"=IFERROR(D{r}/B{r},0)")
        for c in range(1, 8):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=FONT, size=10)
            cell.border = BORDER
            if c in (6, 7):
                cell.number_format = "0%"
    src_last = s0 + 1 + len(SOURCES)

    # Role outcomes
    o0 = src_last + 3
    ws.cell(row=o0, column=1, value="Outcome by role").font = Font(name=FONT, bold=True, size=11, color=NAVY)
    for c, h in enumerate(["Role", "Applicants"] + OUTCOMES, start=1):
        hdr(ws.cell(row=o0 + 1, column=c, value=h))
    for i, role in enumerate(ROLES):
        r = o0 + 2 + i
        ws.cell(row=r, column=1, value=role)
        ws.cell(row=r, column=2, value=f"=COUNTIF({rng('D')},A{r})")
        for j, oc in enumerate(OUTCOMES):
            ws.cell(row=r, column=3 + j, value=f'=COUNTIFS({rng("D")},$A{r},{rng("L")},"{oc}")')
        for c in range(1, 3 + len(OUTCOMES)):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=FONT, size=10)
            cell.border = BORDER

    chart = BarChart()
    chart.type = "col"
    chart.title = "Hiring funnel"
    chart.y_axis.title = "Candidates"
    chart.add_data(Reference(ws, min_col=3, min_row=r0 + 1, max_row=funnel_last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=funnel_first, max_row=funnel_last))
    chart.legend = None
    chart.height, chart.width = 7.5, 13
    ws.add_chart(chart, "I4")

    for col, w in zip("ABCDEFG", [30, 14, 14, 16, 18, 15, 12]):
        ws.column_dimensions[col].width = w
    return ws


def add_readme(wb):
    ws = wb.create_sheet("Read Me", 1)
    ws.sheet_view.showGridLines = False
    lines = [
        ("Recruitment Tracker: how to use", True),
        ("", False),
        ("Candidates sheet", True),
        ("Blue text (columns A to M) is typed by the recruiter. Black text (columns N to P) is calculated. Do not overwrite black cells.", False),
        ("Role, Source, Recruiter, Stage and Outcome use dropdown lists. Edit the allowed values on the Lists sheet.", False),
        ("Stage = the furthest stage the candidate has reached. Outcome = the current status (In Progress, Rejected, Withdrawn, Joined).", False),
        ("Add new candidates in the next empty row. Formulas and dropdowns are ready down to row 500.", False),
        ("Amber Days to Screen means the candidate waited longer than the SLA set on the Dashboard (cell B4).", False),
        ("", False),
        ("Dashboard sheet", True),
        ("Every number is a formula. Nothing needs refreshing. Yellow cell B4 is the only editable assumption.", False),
        ("", False),
        ("About the data", True),
        ("All 120 candidates are synthetic (fictional names, example.com emails). Generated by src/generate_sample_data.py.", False),
    ]
    for r, (text, bold) in enumerate(lines, start=1):
        c = ws.cell(row=r, column=1, value=text or None)
        c.font = Font(name=FONT, bold=bold, size=12 if r == 1 else 10, color=NAVY if bold else "000000")
    ws.column_dimensions["A"].width = 130


def build(csv_path, out_path):
    rows = read_csv(csv_path)
    wb = Workbook()
    wb.remove(wb.active)
    ranges = add_lists_sheet(wb)
    add_candidates_sheet(wb, rows, ranges)
    add_dashboard(wb)
    add_readme(wb)
    wb.move_sheet("Candidates", offset=-(wb.sheetnames.index("Candidates") - 1))
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--input", default="data/sample_candidates.csv")
    p.add_argument("--out", default="output/recruitment_tracker.xlsx")
    a = p.parse_args()
    print(f"Tracker written to {build(a.input, a.out)}")


if __name__ == "__main__":
    main()
